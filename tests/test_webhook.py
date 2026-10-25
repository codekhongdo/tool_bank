from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.api.webhooks.sepay as sepay_module
from app.main import app
from app.services.excel_exporter import ExcelExporter


TOKEN = "test-webhook-token"


def _configure_test_dependencies(monkeypatch, tmp_path: Path) -> Path:
    """Route Excel writes to an isolated temporary workbook."""
    output_file = tmp_path / "payments.xlsx"
    monkeypatch.setattr(
        sepay_module,
        "get_settings",
        lambda: SimpleNamespace(SEPAY_WEBHOOK_AUTH_TOKEN=TOKEN),
    )
    monkeypatch.setattr(
        sepay_module,
        "ExcelExporter",
        lambda: ExcelExporter(str(output_file)),
    )
    return output_file


def _payload(transaction_id: str) -> dict:
    return {
        "content": f"NGUYEN VAN A THCSC3 {transaction_id}",
        "transferAmount": 100000,
        "transactionDate": "2026-10-25 10:00:00",
        "id": transaction_id,
    }


def test_sai_token_va_body_sai_van_bi_401(monkeypatch, tmp_path):
    _configure_test_dependencies(monkeypatch, tmp_path)
    client = TestClient(app)

    response = client.post(
        "/api/webhooks/sepay",
        content="{body khong hop le",
        headers={"Authorization": "Bearer sai-token"},
    )

    assert response.status_code == 401


def test_webhook_dong_thoi_khong_lam_hong_file(monkeypatch, tmp_path):
    output_file = _configure_test_dependencies(monkeypatch, tmp_path)

    def send_request(index: int):
        with TestClient(app) as client:
            return client.post(
                "/api/webhooks/sepay",
                json=_payload(f"tx-concurrent-{index}"),
                headers={"Authorization": f"Bearer {TOKEN}"},
            )

    with ThreadPoolExecutor(max_workers=10) as executor:
        responses = list(executor.map(send_request, range(10)))

    assert all(response.status_code == 200 for response in responses)
    assert all(response.json()["status"] == "ok" for response in responses)

    workbook = load_workbook(output_file)
    worksheet = workbook["payments"]
    assert worksheet.max_row == 11


def test_gui_lai_dong_thoi_cung_id_chi_ghi_1_lan(monkeypatch, tmp_path):
    output_file = _configure_test_dependencies(monkeypatch, tmp_path)

    def send_request(_index: int):
        with TestClient(app) as client:
            return client.post(
                "/api/webhooks/sepay",
                json=_payload("tx-duplicate"),
                headers={"Authorization": f"Bearer {TOKEN}"},
            )

    with ThreadPoolExecutor(max_workers=10) as executor:
        responses = list(executor.map(send_request, range(10)))

    assert all(response.status_code == 200 for response in responses)
    assert sum(response.json()["status"] == "ok" for response in responses) == 1
    assert sum(response.json()["status"] == "duplicate" for response in responses) == 9

    workbook = load_workbook(output_file)
    worksheet = workbook["payments"]
    assert worksheet.max_row == 2
