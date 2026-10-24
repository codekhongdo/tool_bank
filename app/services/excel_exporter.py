from pathlib import Path
from openpyxl import Workbook, load_workbook
from app.models.payment import Payment
import threading

HEADER = ["thoiGian", "ten", "maThamChieu", "soTien", "noiDungGoc", "transactionId"]
SHEET_NAME = "payments"


class ExcelExporter:
    _WRITE_LOCK = threading.Lock()

    def __init__(self, file_path: str = r"d:\tool_bank\data\exports\payments.xlsx"):
        self.file_path = Path(file_path)
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

    def _get_workbook_and_sheet(self):
        if self.file_path.exists():
            wb = load_workbook(self.file_path)
            if SHEET_NAME in wb.sheetnames:
                ws = wb[SHEET_NAME]
            else:
                ws = wb.create_sheet(SHEET_NAME)
                ws.append(HEADER)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET_NAME
            ws.append(HEADER)
        return wb, ws

    def _is_duplicate(self, ws, transactionId: str) -> bool:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[5] == transactionId:
                return True
        return False

    def append(self, payment: Payment) -> bool:
        with self._WRITE_LOCK:
            wb, ws = self._get_workbook_and_sheet()
            if self._is_duplicate(ws, payment.transactionId):
                return False
            ws.append([
                payment.thoiGian,
                payment.ten,
                payment.maThamChieu,
                payment.soTien,
                payment.noiDungGoc,
                payment.transactionId,
            ])
            wb.save(self.file_path)
            return True

